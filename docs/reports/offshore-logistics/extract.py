#!/usr/bin/env python3
"""
Pull the daily narrative out of a DWOS weekly workbook into JSON.

    python3 extract.py <workbook.xlsx> <report-date>   # e.g. 2026-09-03

Writes data/<report-date>.log.json. The workbooks are not committed: they carry
KOC operational data and arrive as email attachments, so they stay outside the
repo and only the extracted log is versioned.

Scope note: this extracts the NARRATIVE only. Trip, truck and fluid counts stay
hand-audited in data/<report-date>.json, because each one is a judgement the
prose does not make for you -- whether "Four Trucks loaded" on Saturday and
"Four Trucks dispatched" on Sunday are four movements or eight, for instance.
Getting that judgement wrong is exactly how the 26-08 report published 46 truck
moves for a week that made 41. A regex cannot make that call; a person can, and
the audit arrays in the data file record which way they called it.
"""

import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
DATE_RE = re.compile(r"^(\d{2})\.(\d{2})\.(\d{4})$")
DAY_FIX = {
    "wedeneday": "Wednesday", "wedneday": "Wednesday", "wednesday": "Wednesday",
    "thurseday": "Thursday", "thursday": "Thursday", "thurday": "Thursday",
    "friday": "Friday", "saturday": "Saturday", "sunday": "Sunday",
    "monday": "Monday", "tuesday": "Tuesday",
}


def clean(v):
    """Normalise a cell: NBSP to space, collapse runs, strip."""
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def day_blocks(ws):
    """Row ranges of each day block, taken from the merged Day column.

    Reading the Day/Date/Status values will not do it. In the 03.09 workbook the
    Thursday 27.08 block carries no Day, no Date and no Status, so a value-driven
    walk silently folds its four operations into Wednesday and reports a 7-day
    week for an 8-day period. The merge geometry still marks the block, because
    whoever built the sheet merged the cells before leaving them empty.
    """
    spans = sorted(
        (r.min_row, r.max_row)
        for r in ws.merged_cells.ranges
        if r.min_col == r.max_col == 2 and r.min_row >= 4
    )
    if spans:
        return spans
    # No merges: fall back to starting a block wherever a date appears.
    starts = [c.row for c in ws["C"] if c.row >= 4 and DATE_RE.match(clean(c.value))]
    return [(s, (starts[i + 1] - 1) if i + 1 < len(starts) else ws.max_row)
            for i, s in enumerate(starts)]


def parse_date(txt):
    m = DATE_RE.match(txt)
    if not m:
        return None
    d, mo, y = (int(g) for g in m.groups())
    try:
        return date(y, mo, d)
    except ValueError:
        return None


def read_sheet(ws, lo_date=None, hi_date=None):
    """One entry per day block, with the block's operations and any defects."""
    days = []
    for lo, hi in day_blocks(ws):
        rows = list(ws.iter_rows(min_row=lo, max_row=hi))
        day, dtxt, status = (clean(rows[0][i].value) for i in range(1, 4))
        entry = {
            "day": DAY_FIX.get(day.lower(), day),
            "date": dtxt,
            "status": status,
            "lines": [],
            "defects": [],
        }
        for r in rows:
            desc = clean(r[4].value)
            if not desc:
                continue
            # "No Night shift operation" is a shift marker, not an operation.
            if re.fullmatch(r"no night shift operation\.?", desc, re.I):
                entry["nightShift"] = False
            else:
                entry["lines"].append(desc)
        if not dtxt:
            entry["defects"].append("no Date in the workbook")
        if not day:
            entry["defects"].append("no Day in the workbook")
        if not status:
            entry["defects"].append("no Status in the workbook")
        entry["_rows"] = [lo, hi]
        days.append(entry)

    # Repair, then derive. Three separate defects live in these columns: a wrong
    # year (the 03.09 sheet opens on 26.08.2025), a missing date, and a
    # misspelled or wrong day name. Each is corrected against a source that
    # cannot itself be wrong in the same way -- the year against the period the
    # sheet declares in its own title, the date against its neighbour, the day
    # name against the date -- and each correction is reported, never silent.
    for i, e in enumerate(days):
        d = parse_date(e["date"])

        if d and lo_date and hi_date and not (lo_date <= d <= hi_date):
            for yr in {lo_date.year, hi_date.year}:
                try:
                    cand = d.replace(year=yr)
                except ValueError:
                    continue
                if lo_date <= cand <= hi_date:
                    e["defects"].append(
                        f"date reads {e['date']}, outside the report period; "
                        f"read as {cand.strftime('%d.%m.%Y')}")
                    d = cand
                    e["date"] = cand.strftime("%d.%m.%Y")
                    break
            else:
                e["defects"].append(
                    f"date {e['date']} is outside the report period and could "
                    f"not be reconciled")

        if d is None and i:
            prev = parse_date(days[i - 1]["date"])
            if prev:
                d = prev + timedelta(days=1)
                e["date"] = d.strftime("%d.%m.%Y")
                e["dateInferred"] = True

        if d:
            proper = d.strftime("%A")
            if e["day"] and e["day"] != proper:
                e["defects"].append(f"Day read {e['day']}, the date is a {proper}")
            e["day"] = proper
            e["iso"] = d.isoformat()
    return days


def period(ws):
    m = re.findall(r"(\d{2}\.\d{2}\.\d{4})", clean(ws["B1"].value))
    return {"from": m[0], "to": m[1]} if len(m) >= 2 else {}


def main():
    if len(sys.argv) < 3:
        sys.exit("usage: extract.py <workbook.xlsx> <report-date>")
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    report_date = sys.argv[2]

    out = {"reportDate": report_date, "source": Path(sys.argv[1]).name}
    for key, sheet in (("port", "Port"), ("vessel", "Vessel")):
        ws = wb[sheet]
        per = period(ws)
        lo, hi = parse_date(per.get("from", "")), parse_date(per.get("to", ""))
        out[key] = {**per, "days": read_sheet(ws, lo, hi)}

    dest = HERE / "data" / f"{report_date}.log.json"
    dest.write_text(json.dumps(out, indent=2, ensure_ascii=False), encoding="utf-8")

    p, v = out["port"], out["vessel"]
    span = parse_date(p.get("to", "")), parse_date(p.get("from", ""))
    expect = (span[0] - span[1]).days + 1 if all(span) else None
    print(f"{dest.name}: {p.get('from')} to {p.get('to')}"
          f"{f' ({expect} days)' if expect else ''}")
    for name, sheet in (("port", p), ("vessel", v)):
        print(f"  {name}: {len(sheet['days'])} blocks, "
              f"{sum(len(d['lines']) for d in sheet['days'])} operations"
              + ("" if expect is None or len(sheet["days"]) == expect
                 else f"  <-- expected {expect}"))
        for d in sheet["days"]:
            for msg in d["defects"]:
                print(f"    {d['date'] or '(no date)'}: {msg}")
            if d.get("dateInferred"):
                print(f"    {d['date']}: date inferred from the previous block")


if __name__ == "__main__":
    main()
